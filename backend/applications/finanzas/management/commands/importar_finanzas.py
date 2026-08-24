from decimal import Decimal, InvalidOperation
from datetime import datetime

from django.contrib.auth import get_user_model
from django.core.management.base import (
    BaseCommand,
    CommandError,
)
from django.db import transaction

from openpyxl import load_workbook

from applications.finanzas.models import (
    CategoriaFinanciera,
    CuentaFinanciera,
    MovimientoFinanciero,
)


User = get_user_model()


class Command(BaseCommand):

    help = (
        "Importa movimientos financieros "
        "desde un archivo Excel."
    )


    def add_arguments(
        self,
        parser,
    ):

        parser.add_argument(
            "archivo",
            type=str,
        )

        parser.add_argument(
            "--usuario",
            type=str,
            required=True,
        )


    def handle(
        self,
        *args,
        **options,
    ):

        archivo = options["archivo"]
        username = options["usuario"]


        # ====================================================
        # USUARIO AUDITORIA
        # ====================================================

        try:

            user = User.objects.get(
                username=username
            )

        except User.DoesNotExist:

            raise CommandError(
                f"No existe el usuario "
                f"'{username}'."
            )


        # ====================================================
        # ABRIR EXCEL
        # ====================================================

        try:

            workbook = load_workbook(
                archivo,
                data_only=True,
            )

        except Exception as exc:

            raise CommandError(
                f"No se pudo abrir el archivo: "
                f"{exc}"
            )


        sheet = workbook.active


        # ====================================================
        # LEER ENCABEZADOS
        # ====================================================

        encabezados = {}

        for column, cell in enumerate(
            sheet[1],
            start=1,
        ):

            if cell.value is None:
                continue

            nombre = self.normalizar_encabezado(
                cell.value
            )

            encabezados[
                nombre
            ] = column


        columnas_requeridas = [
            "fecha",
            "descripcion",
            "tipo movimiento",
            "monto",
            "categoria",
            "cuenta origen",
            "cuenta destino",
        ]


        faltantes = [
            columna
            for columna in columnas_requeridas
            if columna not in encabezados
        ]


        if faltantes:

            raise CommandError(
                "Faltan columnas en el Excel: "
                + ", ".join(faltantes)
            )


        # ====================================================
        # CONTADORES
        # ====================================================

        total_movimientos = 0
        total_cuentas_creadas = 0
        total_categorias_creadas = 0
        total_transferencias = 0


        # ====================================================
        # IMPORTACION
        # ====================================================

        with transaction.atomic():

            for row_number in range(
                2,
                sheet.max_row + 1,
            ):

                fila = self.obtener_fila(
                    sheet,
                    row_number,
                    encabezados,
                )


                # Ignorar fila vacía
                if not any(
                    value is not None
                    and str(value).strip()
                    for value in fila.values()
                ):
                    continue


                fecha = self.parse_fecha(
                    fila["fecha"]
                )

                descripcion = self.texto(
                    fila["descripcion"]
                )

                tipo_excel = self.texto(
                    fila["tipo movimiento"]
                )

                monto = self.parse_monto(
                    fila["monto"]
                )

                categoria_nombre = self.texto(
                    fila["categoria"]
                )

                cuenta_origen_nombre = self.texto(
                    fila["cuenta origen"]
                )

                cuenta_destino_nombre = self.texto(
                    fila["cuenta destino"]
                )


                # ============================================
                # TIPO MOVIMIENTO
                # ============================================

                tipo = self.obtener_tipo_movimiento(
                    tipo_excel
                )


                # ============================================
                # CUENTA ORIGEN
                # ============================================

                cuenta_origen = None

                if cuenta_origen_nombre:

                    cuenta_origen, creada = (
                        self.obtener_o_crear_cuenta(
                            nombre=cuenta_origen_nombre,
                            user=user,
                        )
                    )

                    if creada:
                        total_cuentas_creadas += 1


                # ============================================
                # CUENTA DESTINO
                # ============================================

                cuenta_destino = None

                if cuenta_destino_nombre:

                    cuenta_destino, creada = (
                        self.obtener_o_crear_cuenta(
                            nombre=cuenta_destino_nombre,
                            user=user,
                        )
                    )

                    if creada:
                        total_cuentas_creadas += 1


                # ============================================
                # CATEGORIA
                # ============================================

                categoria = None


                if (
                    tipo
                    != MovimientoFinanciero.Tipo.TRANSFERENCIA
                ):

                    if categoria_nombre:

                        categoria, creada = (
                            self.obtener_o_crear_categoria(
                                nombre=categoria_nombre,
                                user=user,
                            )
                        )

                        if creada:
                            total_categorias_creadas += 1


                else:

                    total_transferencias += 1


                # ============================================
                # VALIDACIONES BASICAS
                # ============================================

                if (
                    tipo
                    == MovimientoFinanciero.Tipo.INGRESO
                    and cuenta_destino is None
                ):

                    raise CommandError(
                        f"Fila {row_number}: "
                        f"el ingreso no tiene "
                        f"Cuenta Destino."
                    )


                if (
                    tipo
                    == MovimientoFinanciero.Tipo.GASTO
                    and cuenta_origen is None
                ):

                    raise CommandError(
                        f"Fila {row_number}: "
                        f"el gasto no tiene "
                        f"Cuenta Origen."
                    )


                if (
                    tipo
                    == MovimientoFinanciero.Tipo.TRANSFERENCIA
                    and (
                        cuenta_origen is None
                        or cuenta_destino is None
                    )
                ):

                    raise CommandError(
                        f"Fila {row_number}: "
                        f"la transferencia necesita "
                        f"Cuenta Origen y Cuenta Destino."
                    )


                if (
                    tipo
                    == MovimientoFinanciero.Tipo.TRANSFERENCIA
                    and cuenta_origen.id
                    == cuenta_destino.id
                ):

                    raise CommandError(
                        f"Fila {row_number}: "
                        f"la cuenta origen y destino "
                        f"son iguales."
                    )


                # ============================================
                # EVITAR DUPLICADOS
                # ============================================

                existe = (
                    MovimientoFinanciero.objects
                    .filter(
                        is_deleted=False,
                        fecha=fecha,
                        descripcion=descripcion,
                        tipo=tipo,
                        monto=monto,
                        categoria=categoria,
                        cuenta_origen=cuenta_origen,
                        cuenta_destino=cuenta_destino,
                    )
                    .exists()
                )


                if existe:

                    self.stdout.write(
                        self.style.WARNING(
                            f"Fila {row_number}: "
                            f"movimiento duplicado, "
                            f"omitido."
                        )
                    )

                    continue


                # ============================================
                # CREAR MOVIMIENTO
                # ============================================

                movimiento = (
                    MovimientoFinanciero.objects.create(
                        fecha=fecha,
                        descripcion=descripcion,
                        tipo=tipo,
                        monto=monto,
                        categoria=categoria,
                        cuenta_origen=cuenta_origen,
                        cuenta_destino=cuenta_destino,
                        observacion="Importado desde Excel",
                        user_made=user,
                        user_updated=user,
                    )
                )


                total_movimientos += 1


                self.stdout.write(
                    f"Fila {row_number}: "
                    f"{movimiento.fecha} | "
                    f"{movimiento.tipo} | "
                    f"{movimiento.descripcion} | "
                    f"${movimiento.monto}"
                )


        # ====================================================
        # RESULTADO
        # ====================================================

        self.stdout.write("")

        self.stdout.write(
            self.style.SUCCESS(
                "IMPORTACION FINALIZADA"
            )
        )

        self.stdout.write(
            f"Movimientos creados: "
            f"{total_movimientos}"
        )

        self.stdout.write(
            f"Cuentas creadas: "
            f"{total_cuentas_creadas}"
        )

        self.stdout.write(
            f"Categorias creadas: "
            f"{total_categorias_creadas}"
        )

        self.stdout.write(
            f"Transferencias: "
            f"{total_transferencias}"
        )


    # ========================================================
    # TIPO MOVIMIENTO
    # ========================================================

    def obtener_tipo_movimiento(
        self,
        value,
    ):

        value = (
            self.texto(value)
            .strip()
            .upper()
        )


        if value in [
            "INGRESO",
            "INGRESOS",
        ]:

            return (
                MovimientoFinanciero
                .Tipo
                .INGRESO
            )


        if value in [
            "GASTO",
            "GASTOS",
        ]:

            return (
                MovimientoFinanciero
                .Tipo
                .GASTO
            )


        if value in [
            "TRANSFERENCIA",
            "TRANSFERENCIAS",
        ]:

            return (
                MovimientoFinanciero
                .Tipo
                .TRANSFERENCIA
            )


        raise ValueError(
            f"Tipo de movimiento "
            f"desconocido: '{value}'"
        )


    # ========================================================
    # CUENTA
    # ========================================================

    def obtener_o_crear_cuenta(
        self,
        nombre,
        user,
    ):

        nombre = nombre.strip()


        cuenta = (
            CuentaFinanciera.objects
            .filter(
                nombre__iexact=nombre,
                is_deleted=False,
            )
            .first()
        )


        if cuenta:

            return cuenta, False


        cuenta = (
            CuentaFinanciera.objects.create(
                nombre=nombre,

                tipo=(
                    self.detectar_tipo_cuenta(
                        nombre
                    )
                ),

                saldo_inicial=Decimal(
                    "0.00"
                ),

                descripcion=(
                    "Importada desde Excel"
                ),

                activa=True,

                user_made=user,
                user_updated=user,
            )
        )


        self.stdout.write(
            self.style.SUCCESS(
                f"  + Cuenta creada: "
                f"{nombre}"
            )
        )


        return cuenta, True


    # ========================================================
    # CATEGORIA
    # ========================================================

    def obtener_o_crear_categoria(
        self,
        nombre,
        user,
    ):

        nombre = nombre.strip()


        categoria = (
            CategoriaFinanciera.objects
            .filter(
                nombre__iexact=nombre,
                is_deleted=False,
            )
            .first()
        )


        if categoria:

            return categoria, False


        categoria = (
            CategoriaFinanciera.objects.create(
                nombre=nombre,

                grupo=None,

                descripcion=(
                    "Importada desde Excel"
                ),

                activo=True,

                user_made=user,
                user_updated=user,
            )
        )


        self.stdout.write(
            self.style.SUCCESS(
                f"  + Categoría creada: "
                f"{nombre}"
            )
        )


        return categoria, True


    # ========================================================
    # TIPO CUENTA
    # ========================================================

    def detectar_tipo_cuenta(
        self,
        nombre,
    ):

        nombre_lower = (
            nombre
            .strip()
            .lower()
        )


        if (
            "efectivo"
            in nombre_lower
        ):

            return (
                CuentaFinanciera
                .Tipo
                .EFECTIVO
            )


        if (
            "naranja"
            in nombre_lower
            or
            "mercado pago"
            in nombre_lower
            or
            "uala"
            in nombre_lower
            or
            "ualá"
            in nombre_lower
        ):

            return (
                CuentaFinanciera
                .Tipo
                .BILLETERA
            )


        return (
            CuentaFinanciera
            .Tipo
            .OTRO
        )


    # ========================================================
    # FILA
    # ========================================================

    def obtener_fila(
        self,
        sheet,
        row_number,
        encabezados,
    ):

        return {
            nombre: sheet.cell(
                row=row_number,
                column=column,
            ).value
            for nombre, column
            in encabezados.items()
        }


    # ========================================================
    # ENCABEZADO
    # ========================================================

    def normalizar_encabezado(
        self,
        value,
    ):

        return (
            str(value)
            .strip()
            .lower()
        )


    # ========================================================
    # TEXTO
    # ========================================================

    def texto(
        self,
        value,
    ):

        if value is None:
            return ""

        return str(
            value
        ).strip()


    # ========================================================
    # MONTO
    # ========================================================

    def parse_monto(
        self,
        value,
    ):

        if value is None:

            raise ValueError(
                "Monto vacío."
            )


        if isinstance(
            value,
            (int, float, Decimal),
        ):

            monto = Decimal(
                str(value)
            )

        else:

            text = (
                str(value)
                .strip()
                .replace("$", "")
                .replace(" ", "")
            )


            if (
                "," in text
                and "." in text
            ):

                # Ej:
                # 1.234.567,89

                if (
                    text.rfind(",")
                    >
                    text.rfind(".")
                ):

                    text = (
                        text
                        .replace(".", "")
                        .replace(",", ".")
                    )

                # Ej:
                # 1,234,567.89

                else:

                    text = (
                        text
                        .replace(",", "")
                    )


            elif "," in text:

                text = (
                    text
                    .replace(",", ".")
                )


            try:

                monto = Decimal(
                    text
                )

            except InvalidOperation:

                raise ValueError(
                    f"Monto inválido: "
                    f"{value}"
                )


        monto = monto.copy_abs()


        if monto <= Decimal(
            "0.00"
        ):

            raise ValueError(
                f"Monto inválido: "
                f"{monto}"
            )


        return monto


    # ========================================================
    # FECHA
    # ========================================================

    def parse_fecha(
        self,
        value,
    ):

        if value is None:

            raise ValueError(
                "Fecha vacía."
            )


        if isinstance(
            value,
            datetime,
        ):

            return value.date()


        if hasattr(
            value,
            "year",
        ):

            return value


        text = (
            str(value)
            .strip()
        )


        formatos = [
            "%d/%m/%Y",
            "%Y-%m-%d",
            "%d-%m-%Y",
        ]


        for formato in formatos:

            try:

                return datetime.strptime(
                    text,
                    formato,
                ).date()

            except ValueError:
                continue


        raise ValueError(
            f"Fecha inválida: "
            f"{value}"
        )